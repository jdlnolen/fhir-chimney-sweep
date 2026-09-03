#!/usr/bin/env python3
"""Validate an agent-authored review and create matching Markdown/Word reports.

This is a report compiler, not a FHIR validator or an automated reviewer.
"""
import argparse
from collections import Counter
from datetime import date
import hashlib
import html
import json
import os
from pathlib import Path
import re
import tempfile
from urllib.parse import urlsplit

AREAS = ('Documentation', 'Examples', 'Module', 'Cross-layer')
KINDS = ('Correction', 'Suggestion', 'Question', 'Missing example')
CHECKS = ('FHIR validator', 'Publisher', 'Terminology', 'Semantic review', 'DOCX visual QA')
META = ('resource', 'fhir_version', 'package', 'build_url', 'source_revision',
        'review_date', 'scope', 'summary', 'recommendation')
FINDING_TEXT = ('id', 'area', 'kind', 'priority', 'title', 'location', 'evidence',
                'proposed_change', 'replacement', 'rationale', 'validation', 'decision')


def record(obj, fields, label):
    if not isinstance(obj, dict) or set(obj) != set(fields):
        raise ValueError(f'{label}: expected exactly {", ".join(fields)}')


def strings(values, label):
    if not isinstance(values, list) or any(not isinstance(s, str) or not s.strip() for s in values):
        raise ValueError(f'{label}: expected a list of nonempty strings')


def texts(obj, fields):
    for field in fields:
        if not isinstance(obj[field], str) or not obj[field].strip():
            raise ValueError(f'{field}: expected a nonempty string')
        if any(ord(c) < 32 and c not in '\n\r\t' for c in obj[field]):
            raise ValueError(f'{field}: XML-incompatible control character')


def url(value):
    parts = urlsplit(value)
    if parts.scheme not in ('http', 'https') or not parts.netloc or re.search(r'[\s<>"\\]', value):
        raise ValueError(f'Expected an HTTP(S) source URL: {value}')


def unique(rows, key):
    result = set()
    for row in rows:
        value = row[key]
        if value in result:
            raise ValueError(f'Duplicate {key}: {value}')
        result.add(value)
    return result


def validate_report(data):
    record(data, META + ('limitations', 'sources', 'inventory', 'findings', 'checks', 'acceptance'), 'review')
    texts(data, META)
    if not re.fullmatch(r'[A-Za-z][A-Za-z0-9]{0,95}', data['resource']):
        raise ValueError('resource: expected a resource name, not a path')
    if not re.fullmatch(r'\d{4}-\d{2}-\d{2}', data['review_date']):
        raise ValueError('review_date must be YYYY-MM-DD')
    date.fromisoformat(data['review_date'])
    url(data['build_url'])
    strings(data['limitations'], 'limitations')
    strings(data['acceptance'], 'acceptance')
    if not data['acceptance']:
        raise ValueError('acceptance must contain final-revision checks')
    for key in ('sources', 'inventory', 'findings', 'checks'):
        if not isinstance(data[key], list):
            raise ValueError(f'{key}: expected a list')
    for row in data['sources']:
        record(row, ('id', 'title', 'url', 'version', 'locator'), 'source')
        texts(row, row.keys())
        url(row['url'])
    sources = unique(data['sources'], 'id')
    if not sources:
        raise ValueError('sources must document the review evidence')
    for row in data['findings']:
        record(row, FINDING_TEXT + ('source_ids', 'dependencies'), 'finding')
        texts(row, FINDING_TEXT)
        strings(row['source_ids'], 'source_ids')
        strings(row['dependencies'], 'dependencies')
        if not row['source_ids'] or not set(row['source_ids']) <= sources:
            raise ValueError('Unknown or absent finding source IDs')
        if row['area'] not in AREAS or row['kind'] not in KINDS or row['priority'] not in ('P1', 'P2', 'P3'):
            raise ValueError('Invalid finding area, kind, or priority')
        if row['priority'] == 'P1' and row['kind'] != 'Correction':
            raise ValueError('P1 is reserved for demonstrated Corrections')
        if row['kind'] == 'Missing example' and row['area'] != 'Examples':
            raise ValueError('Missing example findings belong to Examples')
    findings = unique(data['findings'], 'id')
    graph = {row['id']: row['dependencies'] for row in data['findings']}
    visited, active = set(), set()

    def visit(node):
        if node not in findings:
            raise ValueError(f'Unknown dependency: {node}')
        if node in active:
            raise ValueError('Finding dependency cycle')
        if node in visited:
            return
        active.add(node)
        for child in graph[node]:
            visit(child)
        active.remove(node)
        visited.add(node)

    for node in graph:
        visit(node)
    for row in data['inventory']:
        record(row, ('area', 'artifact', 'status', 'note', 'source_ids', 'finding_ids'), 'inventory')
        texts(row, ('area', 'artifact', 'status', 'note'))
        for key, allowed in [('source_ids', sources), ('finding_ids', findings)]:
            strings(row[key], key)
            if not set(row[key]) <= allowed:
                raise ValueError(f'Unknown inventory {key}')
        if row['area'] not in AREAS or row['status'] not in ('reviewed', 'not-reviewed', 'excluded'):
            raise ValueError('Invalid inventory area or status')
        if row['status'] == 'reviewed' and not row['source_ids']:
            raise ValueError('Reviewed inventory entries need evidence sources')
    areas = {row['area'] for row in data['inventory']}
    for area in AREAS[:3]:
        if area not in areas:
            raise ValueError(f'Inventory must account for {area}')
    for row in data['checks']:
        record(row, ('name', 'status', 'details'), 'check')
        texts(row, row.keys())
        if row['status'] not in ('passed', 'failed', 'not-run', 'not-applicable'):
            raise ValueError('Invalid check status')
    if not set(CHECKS) <= unique(data['checks'], 'name'):
        raise ValueError(f'checks must include {", ".join(CHECKS)}')
    recommendation = data['recommendation']
    if recommendation not in ('Changes required', 'No material inconsistencies identified', 'Incomplete review'):
        raise ValueError('Invalid recommendation')
    if any(row['status'] == 'not-reviewed' for row in data['inventory']) and recommendation != 'Incomplete review':
        raise ValueError('Unreviewed inventory requires Incomplete review')
    if recommendation == 'No material inconsistencies identified':
        if any(row['kind'] in ('Correction', 'Question') for row in data['findings']) or any(row['status'] == 'failed' for row in data['checks']):
            raise ValueError('Clean recommendation conflicts with findings or failed checks')
        if next(row for row in data['checks'] if row['name'] == 'Semantic review')['status'] != 'passed':
            raise ValueError('Clean recommendation requires a completed semantic review')


def content_blocks(data):
    validate_report(data)
    blocks = []

    def add(kind, value):
        blocks.append((kind, value))

    digest = hashlib.sha256(json.dumps(data, sort_keys=True, ensure_ascii=False).encode()).hexdigest()[:16]
    add('title', f'FHIR {data["resource"]}: Chimney Sweep')
    add('p', 'Resource continuity and prepublication review')
    for label, key in [('Recommendation', 'recommendation'), ('Reviewed', 'review_date'),
                       ('FHIR version', 'fhir_version'), ('Package', 'package'), ('Source revision', 'source_revision')]:
        add('meta', f'{label}: {data[key]}')
    add('link', ('Build', data['build_url']))
    add('meta', f'Report content ID: {digest} (identical in Markdown and Word)')
    add('p', data['summary'])
    counts = Counter(f['priority'] for f in data['findings'])
    kinds = Counter(f['kind'] for f in data['findings'])
    add('p', 'Triage: ' + '; '.join(f'{k}: {counts[k]}' for k in ('P1', 'P2', 'P3')) + '. ' +
        '; '.join(f'{k}: {kinds[k]}' for k in KINDS) + '.')
    add('h1', 'Scope and review boundary')
    add('p', data['scope'])
    add('p', 'Advisory change list only. No FHIR source was modified by this sweep. '
        'Parsing, FHIR validation, terminology verification, semantic review, and document rendering are distinct checks. '
        'This report does not confer clinical or publication approval.')
    add('h1', 'Review inventory')
    add('table', [('Area', 'Artifact', 'Disposition', 'Finding IDs')] + [
        (r['area'], r['artifact'], r['status'], ', '.join(r['finding_ids']) or 'None') for r in data['inventory']])
    for row in data['inventory']:
        add('meta', f'{row["area"]} / {row["artifact"]}: {row["note"]} Sources: {", ".join(row["source_ids"]) or "None available"}.')
    sections = [(f'{a} findings', [f for f in data['findings'] if f['area'] == a and f['kind'] != 'Missing example']) for a in AREAS]
    sections.append(('Suggested missing examples', [f for f in data['findings'] if f['kind'] == 'Missing example']))
    for title, rows in sections:
        add('h1', title)
        if not rows:
            add('p', 'No findings recorded for this section. See inventory and limitations for actual coverage.')
        for f in sorted(rows, key=lambda item: item['priority']):
            add('h2', f'{f["id"]} | {f["priority"]} | {f["kind"]}: {f["title"]}')
            for label, key in [('Location', 'location'), ('Evidence', 'evidence')]:
                add('p', f'{label}: {f[key]}')
            add('meta', 'Evidence sources: ' + ', '.join(f['source_ids']))
            add('p', 'Proposed change: ' + f['proposed_change'])
            add('p', 'Replacement / minimum content:')
            add('code', f['replacement'])
            for label, key in [('Rationale', 'rationale'), ('Acceptance check', 'validation'), ('Human decision', 'decision')]:
                add('p', f'{label}: {f[key]}')
            add('meta', 'Depends on: ' + (', '.join(f['dependencies']) or 'None'))
    add('h1', 'Verification performed')
    for row in data['checks']:
        add('p', f'{row["name"]}: {row["status"]}. {row["details"]}')
    add('h1', 'Publication acceptance checklist')
    for item in data['acceptance']:
        add('bullet', 'Pending: ' + item)
    add('h1', 'Limitations and decision gates')
    for item in data['limitations'] or ['No additional limitations recorded. Check the verification section before relying on this review.']:
        add('bullet', item)
    for f in data['findings']:
        if f['kind'] == 'Question':
            add('p', f'{f["id"]}: {f["decision"]}')
    add('h1', 'Sources')
    for source in data['sources']:
        add('link', (f'{source["id"]}: {source["title"]}', source['url']))
        add('meta', f'Version/snapshot: {source["version"]}. Location: {source["locator"]}.')
    return blocks


def escape_md(value):
    value = html.escape(value, quote=False)
    return re.sub(r'([\\`*_{}\[\]()#+!|>~])', r'\\\1', value).replace('\n', '<br>')


def render_markdown(blocks):
    result = []
    for kind, value in blocks:
        if kind == 'table':
            for i, row in enumerate(value):
                result.append('| ' + ' | '.join(escape_md(c) for c in row) + ' |')
                if i == 0:
                    result.append('| ' + ' | '.join('---' for _ in row) + ' |')
        elif kind == 'link':
            result.append(f'[{escape_md(value[0])}](<{value[1]}>)')
        elif kind == 'code':
            fence = '`' * max(3, max((len(s) + 1 for s in re.findall(r'`+', value)), default=3))
            result.append(f'{fence}\n{value}\n{fence}')
        else:
            prefix = {'title': '# ', 'h1': '## ', 'h2': '### ', 'bullet': '- '}.get(kind, '')
            result.append(prefix + escape_md(value))
        result.append('')
    return '\n'.join(result)


def visible_texts(blocks):
    values = []
    for kind, value in blocks:
        chunks = [c for row in value for c in row] if kind == 'table' else ([value[0]] if kind == 'link' else [value])
        values.extend(s for chunk in chunks for s in re.split(r'[\r\n\t]', chunk) if s)
    return values


def render_docx(blocks, path):
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.shared import Inches, Pt, RGBColor, Twips

    def element(tag, **attrs):
        node = OxmlElement('w:' + tag)
        for k, v in attrs.items():
            node.set(qn('w:' + k), str(v))
        return node

    doc = Document()
    doc.core_properties.author = 'FHIR-chimney-sweep'
    doc.core_properties.title = blocks[0][1]
    doc.core_properties.subject = 'Advisory FHIR resource continuity review'
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(.492)
    styles = {
        'Normal': (11, '000000', 0, 6, 1.25, 'Calibri'),
        'Title': (23, '0B2545', 0, 8, 1.1, 'Calibri'),
        'Heading 1': (16, '2E74B5', 18, 10, 1.25, 'Calibri'),
        'Heading 2': (13, '2E74B5', 14, 7, 1.25, 'Calibri'),
        'Heading 3': (12, '1F4D78', 10, 5, 1.25, 'Calibri'),
        'Sweep Metadata': (9, '555555', 0, 4, 1.1, 'Calibri'),
        'Sweep Code': (9, '000000', 0, 4, 1, 'Courier New'),
        'Sweep Table': (9, '000000', 0, 4, 1.1, 'Calibri'),
        'Sweep Bullet': (11, '000000', 0, 4, 1.25, 'Calibri'),
    }
    for name, (size, color, before, after, spacing, font) in styles.items():
        s = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        s.font.name, s.font.size, s.font.color.rgb = font, Pt(size), RGBColor.from_string(color)
        s.font.bold = name.startswith('Heading')
        p = s.paragraph_format
        p.space_before, p.space_after, p.line_spacing = Pt(before), Pt(after), spacing
        p.keep_with_next = name.startswith('Heading') or name == 'Title'
        p.widow_control = True
    title_properties = doc.styles['Title'].element.find(qn('w:pPr'))
    for border in list(title_properties.findall(qn('w:pBdr'))):
        title_properties.remove(border)
    numbering = doc.part.numbering_part.element
    abstract_id = max(int(n.get(qn('w:abstractNumId'))) for n in numbering.findall(qn('w:abstractNum'))) + 1
    num_id = max(int(n.get(qn('w:numId'))) for n in numbering.findall(qn('w:num'))) + 1
    abstract = element('abstractNum', abstractNumId=abstract_id)
    lvl = element('lvl', ilvl=0)
    for node in [element('start', val=1), element('numFmt', val='bullet'), element('lvlText', val='•'), element('lvlJc', val='left')]:
        lvl.append(node)
    pp = element('pPr')
    pp.append(element('ind', left=540, hanging=271))
    tabs = element('tabs')
    tabs.append(element('tab', val='num', pos=540))
    pp.append(tabs)
    lvl.append(pp)
    abstract.append(lvl)
    numbering.append(abstract)
    num = element('num', numId=num_id)
    num.append(element('abstractNumId', val=abstract_id))
    numbering.append(num)
    section.header.paragraphs[0].text = 'FHIR-chimney-sweep | Resource continuity review'
    section.header.paragraphs[0].style = doc.styles['Sweep Metadata']
    footer = section.footer.paragraphs[0]
    footer.style = doc.styles['Sweep Metadata']
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Page ')
    field = element('fldSimple', instr='PAGE')
    footer._p.append(field)

    for kind, value in blocks:
        if kind == 'table':
            table = doc.add_table(rows=len(value), cols=4)
            table.autofit = False
            widths = [1500, 3600, 1500, 2760]
            table._tbl.tblPr.find(qn('w:tblW')).set(qn('w:w'), '9360')
            table._tbl.tblPr.find(qn('w:tblW')).set(qn('w:type'), 'dxa')
            table._tbl.tblPr.append(element('tblInd', w=120, type='dxa'))
            margins = element('tblCellMar')
            for side, width in [('top', 80), ('bottom', 80), ('start', 120), ('end', 120)]:
                margins.append(element(side, w=width, type='dxa'))
            table._tbl.tblPr.append(margins)
            borders = element('tblBorders')
            for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
                borders.append(element(side, val='single', sz=4, color='CBD5E1'))
            table._tbl.tblPr.append(borders)
            for col, width in zip(table.columns, widths):
                col.width = Twips(width)
            for i, values in enumerate(value):
                for cell, width, text in zip(table.rows[i].cells, widths, values):
                    cell.width = Twips(width)
                    cell.text = text
                    for p in cell.paragraphs:
                        p.style = doc.styles['Sweep Table']
                    cell._tc.get_or_add_tcPr().append(element('vAlign', val='center'))
                    if i == 0:
                        cell._tc.get_or_add_tcPr().append(element('shd', fill='E8EEF5'))
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
                if i == 0:
                    table.rows[i]._tr.get_or_add_trPr().append(element('tblHeader'))
        elif kind == 'link':
            p = doc.add_paragraph(style='Sweep Metadata')
            p.paragraph_format.keep_with_next = True
            link = OxmlElement('w:hyperlink')
            link.set(qn('r:id'), doc.part.relate_to(value[1], RT.HYPERLINK, is_external=True))
            run = element('r')
            props = element('rPr')
            props.append(element('color', val='2E74B5'))
            props.append(element('u', val='single'))
            run.append(props)
            t = element('t')
            t.text = value[0]
            run.append(t)
            link.append(run)
            p._p.append(link)
        else:
            style = {'title': 'Title', 'h1': 'Heading 1', 'h2': 'Heading 2', 'meta': 'Sweep Metadata', 'code': 'Sweep Code', 'bullet': 'Sweep Bullet'}.get(kind, 'Normal')
            p = doc.add_paragraph(value, style=style)
            if kind == 'bullet':
                np = element('numPr')
                np.append(element('ilvl', val=0))
                np.append(element('numId', val=num_id))
                p._p.get_or_add_pPr().append(np)
    doc.save(path)


def write_reports(data, out_dir, overwrite=False):
    blocks = content_blocks(data)
    out_dir = Path(out_dir)
    stem = f'FHIR-{data["resource"]}-Chimney-Sweep-{data["review_date"]}'
    targets = [out_dir / (stem + ext) for ext in ('.md', '.docx')]
    for target in targets:
        if target.is_symlink() or (target.exists() and (not overwrite or not target.is_file())):
            raise FileExistsError(f'Refusing to replace {target}; choose a fresh directory or --overwrite for existing regular report files')
    out_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='.sweep-', dir=out_dir) as staging:
        temp_md, temp_docx = [Path(staging) / p.name for p in targets]
        temp_md.write_text(render_markdown(blocks), encoding='utf-8')
        render_docx(blocks, temp_docx)
        published = []
        try:
            for source, target in zip((temp_md, temp_docx), targets):
                if overwrite:
                    source.replace(target)
                else:
                    # Atomic no-clobber publication, even if another writer raced us.
                    os.link(source, target)
                    published.append((source, target))
        except OSError:
            for source, target in published:
                if target.exists() and not target.is_symlink() and source.samefile(target):
                    target.unlink()
            raise
    return tuple(targets)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('review', type=Path, nargs='?')
    parser.add_argument('--out-dir', type=Path)
    parser.add_argument('--overwrite', action='store_true')
    parser.add_argument('--check', action='store_true', help='Validate review data without generating reports')
    parser.add_argument('--example', action='store_true', help='Print synthetic input format demonstration')
    args = parser.parse_args()
    if args.example:
        print((Path(__file__).resolve().parents[1] / 'assets/example-review.json').read_text())
        return
    if not args.review or (not args.check and not args.out_dir):
        parser.error('supply review.json and --out-dir, or review.json --check')
    try:
        data = json.loads(args.review.read_text(encoding='utf-8'))
        if args.check:
            validate_report(data)
            print('Review structure and cross-references passed. Evidence truth and completeness require human review.')
        else:
            for path in write_reports(data, args.out_dir, args.overwrite):
                print(path.resolve())
            print('Generated both formats. DOCX visual QA still requires rendering and page inspection.')
    except (ValueError, OSError) as exc:
        parser.exit(1, f'Report error: {exc}\n')


if __name__ == '__main__':
    main()
