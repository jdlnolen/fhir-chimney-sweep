import hashlib
from pathlib import Path
import re
import unittest
from zipfile import ZipFile

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / 'plugins/fhir-chimney-sweep/skills/fhir-chimney-sweep/references/sample-reports'


class SampleReportTests(unittest.TestCase):
    def test_dated_samples_preserve_both_report_files(self):
        samples = sorted(SAMPLES.glob('*/SHA256SUMS'))
        self.assertTrue(samples, 'At least one dated sample must be bundled')
        for manifest in samples:
            with self.subTest(sample=manifest.parent.name):
                entries = [line.split(maxsplit=1) for line in manifest.read_text().splitlines()]
                self.assertEqual(len(entries), 2)
                self.assertEqual({Path(name).suffix for _, name in entries}, {'.md', '.docx'})
                for expected, name in entries:
                    self.assertEqual(Path(name).name, name, 'Sample filenames must be local')
                    data = (manifest.parent / name).read_bytes()
                    self.assertEqual(hashlib.sha256(data).hexdigest(), expected, name)
                self.assertEqual(len({Path(name).stem for _, name in entries}), 1)

    def test_samples_have_matching_content_ids_and_valid_word_packages(self):
        reports = sorted(SAMPLES.glob('*/*.docx'))
        self.assertTrue(reports)
        for word_path in reports:
            with self.subTest(sample=word_path.parent.name):
                markdown = word_path.with_suffix('.md').read_text()
                md_id = re.search(r'Report content ID: ([0-9a-f]{16})\b', markdown)
                self.assertIsNotNone(md_id)
                with ZipFile(word_path) as package:
                    self.assertIsNone(package.testzip())
                    self.assertIn('word/document.xml', package.namelist())
                text = '\n'.join(p.text for p in Document(word_path).paragraphs)
                word_id = re.search(r'Report content ID: ([0-9a-f]{16})\b', text)
                self.assertIsNotNone(word_id)
                self.assertEqual(md_id.group(1), word_id.group(1))

    def test_observation_sample_uses_classification_area_numeric_order(self):
        sample = SAMPLES / 'observation-2026-09-03/FHIR-Observation-Chimney-Sweep-2026-09-03.md'
        lines = sample.read_text().splitlines()
        classes = [
            'Must fix (implementation or testing impact)',
            'Minor fix',
            'Net new addition',
        ]
        seen_classes = []
        grouped_ids = {}
        current_class = current_area = None
        for line in lines:
            heading = line[3:].replace('\\(', '(').replace('\\)', ')') if line.startswith('## ') else ''
            if heading in classes:
                current_class = heading
                current_area = None
                seen_classes.append(current_class)
            elif current_class and line.startswith('### '):
                current_area = line[4:]
            elif current_class and current_area and line.startswith('#### '):
                match = re.match(r'#### [A-Z-]+?(\d+)\b', line)
                self.assertIsNotNone(match, line)
                grouped_ids.setdefault((current_class, current_area), []).append(int(match.group(1)))
        self.assertEqual(seen_classes, classes)
        self.assertEqual(sum(len(ids) for ids in grouped_ids.values()), 54)
        for group, ids in grouped_ids.items():
            self.assertEqual(ids, sorted(ids), group)

        word = Document(sample.with_suffix('.docx'))
        finding_paragraphs = [p for p in word.paragraphs if re.match(r'^[A-Z-]+\d+ \| P[123] \|', p.text)]
        self.assertEqual(len(finding_paragraphs), 54)
        self.assertTrue(all(p.style.name == 'Heading 3' for p in finding_paragraphs))


if __name__ == '__main__':
    unittest.main()
